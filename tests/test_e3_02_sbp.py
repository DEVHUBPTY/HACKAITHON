"""E3-02 · Fuente D: series agregadas de la SBP (conversión, validación, manifest, catálogo y extracción sin red)."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

import pytest

from scripts import catalogo, conversion, extraer, sbp, validar_snapshot
from tests.conftest import escribir_manifest, escribir_sbp_sintetico

SERIES = ["SBP-MOROSIDAD-SISTEMA", "SBP-MOROSOS-SISTEMA", "SBP-PROVISIONES-SISTEMA"]


def _filas(data: Path) -> list[dict[str, str]]:
    with (data / "processed" / sbp.NOMBRE_CSV).open(encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def test_el_csv_tiene_las_columnas_de_la_spec_y_12_periodos_por_serie(data_sintetica: Path) -> None:
    filas = _filas(data_sintetica)
    assert list(filas[0]) == sbp.COLUMNAS_SBP
    for serie in SERIES:
        periodos = sorted(f["periodo"] for f in filas if f["id_serie"] == serie)
        assert periodos == [f"2024-{m:02d}" for m in range(1, 13)], serie


def test_cada_fila_lleva_unidad_pagina_informe_url_y_condiciones(data_sintetica: Path) -> None:
    for f in _filas(data_sintetica):
        assert f["unidad"] and f["informe"] and f["url"].startswith("https://") and f["fecha_extraccion"].endswith("Z")
        assert f["pagina"].startswith("hoja «") and "celda" in f["pagina"]
        assert "Pendiente de verificar" in f["condiciones"] and "autorización" in f["condiciones"]
    assert sbp.id_dato("SBP-MOROSOS-SISTEMA", "2024-03") == "SBP-MOROSOS-SISTEMA-2024-03"


def test_ningun_campo_identifica_clientes_ni_entidades_individuales(data_sintetica: Path) -> None:
    """Solo salen las series agregadas del sistema: el banco individual del informe (fila 7, 7777) no aparece en ninguna celda."""
    filas = _filas(data_sintetica)
    assert {f["id_serie"] for f in filas} == set(SERIES)
    texto = json.dumps(filas, ensure_ascii=False).casefold()
    assert "banco individual" not in texto and "7777" not in texto
    assert not {"cliente", "banco", "entidad", "nombre_banco", "id_cliente"} & {c.casefold() for c in sbp.COLUMNAS_SBP}
    cfg = conversion.cargar_config()["sbp"]
    etiquetas = {s["etiqueta_fila"].casefold() for s in cfg["series"].values()}
    assert etiquetas == {"sistema bancario", "morosos / sistema bancario"}


def test_los_valores_son_los_del_informe_en_su_celda_y_conservan_la_unidad(data_sintetica: Path) -> None:
    filas = {(f["id_serie"], f["periodo"]): f for f in _filas(data_sintetica)}
    enero = filas[("SBP-MOROSOS-SISTEMA", "2024-01")]
    assert float(enero["valor"]) == 1000.0 + 7           # 2023-06 es la columna 2; enero de 2024 es el mes 8 (índice 7)
    assert enero["pagina"] == "hoja «Morosos», celda I6"
    assert enero["unidad"] == "millones de balboas"
    assert float(filas[("SBP-MOROSIDAD-SISTEMA", "2024-01")]["valor"]) == pytest.approx(0.015 + 7 / 10000)


def test_un_valor_vacio_queda_nulo_y_nunca_cero(tmp_path: Path) -> None:
    raw = tmp_path / "raw"
    escribir_sbp_sintetico(raw / "sbp", nulo_en="2024-05")
    filas = conversion.sbp.convertir_sbp(raw, conversion.cargar_config())
    mayo = next(f for f in filas if f["id_serie"] == "SBP-MOROSOS-SISTEMA" and f["periodo"] == "2024-05")
    assert mayo["valor"] is None
    ruta = tmp_path / "s.csv"
    conversion.escribir_csv(ruta, sbp.COLUMNAS_SBP, filas)
    celda = next(f for f in csv.DictReader(ruta.open(encoding="utf-8")) if f["id_serie"] == "SBP-MOROSOS-SISTEMA" and f["periodo"] == "2024-05")
    assert celda["valor"] == ""


def test_si_la_sbp_mueve_la_fila_la_conversion_falla_en_vez_de_leer_otra(tmp_path: Path) -> None:
    import openpyxl

    raw = tmp_path / "raw"
    escribir_sbp_sintetico(raw / "sbp")
    ruta = next((raw / "sbp").glob("sbp_morosos_*.xlsx"))
    libro = openpyxl.load_workbook(ruta)
    libro.active.cell(row=6, column=1, value="BANCO INDIVIDUAL SA")
    libro.save(ruta)
    with pytest.raises(sbp.ErrorSbp, match="se llama"):
        conversion.sbp.convertir_sbp(raw, conversion.cargar_config())


def test_si_faltan_meses_del_rango_la_conversion_falla(tmp_path: Path) -> None:
    import openpyxl

    raw = tmp_path / "raw"
    escribir_sbp_sintetico(raw / "sbp")
    ruta = next((raw / "sbp").glob("sbp_provisiones_*.xlsx"))
    libro = openpyxl.load_workbook(ruta)
    libro.active.cell(row=4, column=12, value="sin fecha")   # enero..: una columna de 2024 deja de ser fecha
    libro.save(ruta)
    with pytest.raises(sbp.ErrorSbp, match="períodos"):
        conversion.sbp.convertir_sbp(raw, conversion.cargar_config())


def test_si_falta_el_crudo_de_un_informe_el_error_lo_nombra(tmp_path: Path) -> None:
    """X90: con un solo informe no se devuelve ``[]`` en silencio: el error nombra el que falta."""
    raw = tmp_path / "raw"
    escribir_sbp_sintetico(raw / "sbp")
    cfg = conversion.cargar_config()
    faltante = next(iter(cfg["sbp"]["informes"]))
    for p in (raw / "sbp").glob(f"*{faltante}*"):
        p.unlink()
    with pytest.raises(sbp.ErrorSbp, match=faltante):
        conversion.sbp.convertir_sbp(raw, cfg)


def test_sin_crudos_de_la_sbp_la_fuente_es_opcional_y_no_toca_el_csv_existente(tmp_path: Path) -> None:
    assert conversion.sbp.convertir_sbp(tmp_path / "raw", conversion.cargar_config()) == []


def test_la_conversion_es_determinista(data_sintetica: Path, config: dict[str, Any]) -> None:
    antes = (data_sintetica / "processed" / sbp.NOMBRE_CSV).read_bytes()
    conversion.convertir_todo(data_sintetica / "raw", data_sintetica / "processed", config)
    assert (data_sintetica / "processed" / sbp.NOMBRE_CSV).read_bytes() == antes


def _estados(informe: dict[str, Any]) -> dict[str, str]:
    return {c["comprobacion"]: c["estado"] for c in informe["comprobaciones"]}


def test_validar_snapshot_acepta_la_fuente_d_y_rechaza_un_periodo_faltante(data_sintetica: Path, config: dict[str, Any]) -> None:
    estados = _estados(validar_snapshot.validar(data_sintetica, config))
    assert estados["sbp:series_y_periodos"] == "ok" and estados["sbp:unidad_pagina_condiciones"] == "ok"
    conversion.escribir_csv(data_sintetica / "processed" / sbp.NOMBRE_CSV, sbp.COLUMNAS_SBP, _filas(data_sintetica)[:-1])
    assert _estados(validar_snapshot.validar(data_sintetica, config))["sbp:series_y_periodos"] == "error"


def test_validar_snapshot_rechaza_una_fila_sin_pagina(data_sintetica: Path, config: dict[str, Any]) -> None:
    filas = _filas(data_sintetica)
    filas[0]["pagina"] = ""
    conversion.escribir_csv(data_sintetica / "processed" / sbp.NOMBRE_CSV, sbp.COLUMNAS_SBP, filas)
    assert _estados(validar_snapshot.validar(data_sintetica, config))["sbp:unidad_pagina_condiciones"] == "error"


def test_validar_snapshot_solo_advierte_si_no_hay_fuente_d(data_sintetica: Path, config: dict[str, Any]) -> None:
    """Un snapshot sin fuente D (sin CSV ni huella en el manifest) solo advierte; el CSV, ya versionado (D-116), no se omite si el manifest lo lista."""
    (data_sintetica / "processed" / sbp.NOMBRE_CSV).unlink()
    ruta = data_sintetica / "manifest.json"
    manifiesto = json.loads(ruta.read_text("utf-8"))
    manifiesto["sha256"].pop("processed/sbp_series.csv")
    ruta.write_text(json.dumps(manifiesto, ensure_ascii=False), encoding="utf-8")
    informe = validar_snapshot.validar(data_sintetica, config)
    assert _estados(informe)["archivo:sbp_series.csv"] == "advertencia"
    assert not [c for c in informe["comprobaciones"] if c["comprobacion"].startswith("sbp:") and c["estado"] == "error"]
    assert _estados(informe).get("manifest:sha256") != "error"


def test_el_manifest_incluye_la_fuente_d_con_su_hash_y_sus_condiciones(data_sintetica: Path) -> None:
    m = json.loads((data_sintetica / "manifest.json").read_text("utf-8"))
    assert m["sha256"]["processed/sbp_series.csv"] == conversion.sha256_archivo(data_sintetica / "processed" / sbp.NOMBRE_CSV)
    assert m["cantidad_por_archivo"]["sbp_series.csv"] == 36
    assert "restringe" in m["licencias"]["sbp"] and "Pendiente de verificar" in m["licencias"]["sbp"]
    assert set(m["consultas"]["sbp"]["series"]) == set(SERIES)
    assert any(k.startswith("raw/sbp/") for k in m["crudos"])
    assert any("SBP" in t for t in m["transformaciones"])


def test_los_crudos_de_la_sbp_no_mueven_la_fecha_de_corte(data_sintetica: Path, config: dict[str, Any]) -> None:
    escribir_sbp_sintetico(data_sintetica / "raw" / "sbp", momento="20271231-235959")
    assert escribir_manifest(data_sintetica, config)["fecha_corte_UTC"] == "2026-10-06T12:05:00Z"


def test_el_catalogo_describe_la_fuente_d_con_sus_huellas(data_sintetica: Path) -> None:
    m = json.loads((data_sintetica / "manifest.json").read_text("utf-8"))
    fila = catalogo.construir_filas(m, data_sintetica / "processed")[3]
    assert fila["Archivo"] == "sbp_series.csv" and fila["Modalidad"] == "Banca"
    assert m["sha256"]["processed/sbp_series.csv"] in fila["SHA-256"]
    assert "sbp_morosos_" in fila["SHA-256"] and "no versionados" in fila["SHA-256"]
    assert "nunca información de clientes" in fila["Cobertura"] and "2026-10-07T12:00:00Z" in fila["Fecha de extracción"]
    assert "Pendiente de verificar" in fila["Licencia / condiciones"]


# ------------------------------------------------------------------ extracción sin red


class _Respuesta:
    def __init__(self, contenido: bytes, texto: str = "") -> None:
        self.content = contenido
        self.text = texto or contenido.decode("utf-8", errors="ignore")
        self.status_code = 200


def test_el_enlace_se_toma_de_la_pagina_sin_el_parametro_de_version() -> None:
    html = '<a href="/documentos/x/2026/08/calidad/Morosos.xlsx?v=4.07">m</a><a href="/documentos/x/Provisiones.xlsx?v=2">p</a>'
    url = extraer._enlace_del_informe(html, "Morosos.xlsx", "https://www.superbancos.gob.pa", "/estadisticas-financieras/cartera-credito")
    assert url == "https://www.superbancos.gob.pa/documentos/x/2026/08/calidad/Morosos.xlsx"
    with pytest.raises(extraer.ErrorDeExtraccion, match="no tiene el enlace"):
        extraer._enlace_del_informe(html, "Otro.xlsx", "https://www.superbancos.gob.pa", "/p")


def test_extraer_sbp_guarda_crudos_inmutables_y_deja_el_registro(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    origen = tmp_path / "origen"
    escribir_sbp_sintetico(origen)
    xlsx = {"Morosos.xlsx": next(origen.glob("sbp_morosos_*")).read_bytes(), "Provisiones.xlsx": next(origen.glob("sbp_provisiones_*")).read_bytes()}
    pagina = '<a href="/documentos/r/Morosos.xlsx?v=1">m</a><a href="/documentos/r/Provisiones.xlsx?v=1">p</a>'
    pedidos: list[str] = []

    def falso_get(url: str, params: Any, config: Any, espera_429: Any = None) -> _Respuesta:
        pedidos.append(url)
        if url.endswith("/robots.txt"):
            return _Respuesta(b"User-agent: *\nDisallow: /admin/\n")
        if url.endswith(".xlsx"):
            return _Respuesta(xlsx[url.rsplit("/", 1)[1]])
        return _Respuesta(pagina.encode())

    monkeypatch.setattr(extraer, "_get", falso_get)
    monkeypatch.setattr(extraer.time, "sleep", lambda s: None)
    raw = tmp_path / "data" / "raw"
    config = conversion.cargar_config()
    guardados = extraer.extraer_sbp(raw, config)
    assert sorted(p.name.split("_")[1] for p in guardados) == ["morosos", "provisiones"]
    assert all(p.read_bytes()[:2] == b"PK" for p in guardados)
    registros = sorted((raw.parent / "registro_extraccion").glob("sbp_*.json"))
    assert len(registros) == 2
    reg = json.loads(registros[0].read_text("utf-8"))
    assert reg["url"].startswith("https://www.superbancos.gob.pa/documentos/r/") and "?" not in reg["url"]
    assert len(reg["sha256"]) == 64 and reg["origen"] and "Pendiente de verificar" in reg["condiciones"]
    assert sbp.convertir_sbp(raw, config)[0]["url"].startswith("https://www.superbancos.gob.pa/documentos/r/")
    with pytest.raises(FileExistsError):
        extraer._guardar(raw / "sbp", guardados[0].name, b"otra cosa")


def test_extraer_sbp_respeta_robots_txt(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    def falso_get(url: str, params: Any, config: Any, espera_429: Any = None) -> _Respuesta:
        return _Respuesta(b"User-agent: *\nDisallow: /estadisticas-financieras/\n") if url.endswith("/robots.txt") else _Respuesta(b"")

    monkeypatch.setattr(extraer, "_get", falso_get)
    with pytest.raises(extraer.ErrorDeExtraccion, match="robots.txt"):
        extraer.extraer_sbp(tmp_path / "raw", conversion.cargar_config())


def test_extraer_sbp_rechaza_algo_que_no_es_un_xlsx(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    def falso_get(url: str, params: Any, config: Any, espera_429: Any = None) -> _Respuesta:
        if url.endswith("/robots.txt"):
            return _Respuesta(b"User-agent: *\n")
        return _Respuesta(b'<a href="/d/Morosos.xlsx">m</a>') if "estadisticas" in url else _Respuesta(b"<html>error</html>")

    monkeypatch.setattr(extraer, "_get", falso_get)
    monkeypatch.setattr(extraer.time, "sleep", lambda s: None)
    with pytest.raises(extraer.ErrorDeExtraccion, match="no devolvió un .xlsx"):
        extraer.extraer_sbp(tmp_path / "raw", conversion.cargar_config())
